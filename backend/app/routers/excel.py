"""Excel (xlsx) import/export for Actions and Users.

Export respects plant scoping; user import/export is admin-only.
Columns are header-name based so users can reorder freely.
"""
import datetime as _dt
import io
import uuid

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.middleware.auth import require_auth, get_current_user
from app.models.models import Action, User, Plant, Department
from app.services.password import hash_password

actions_excel = APIRouter(prefix="/api/actions/excel", tags=["Excel"], dependencies=[Depends(require_auth)])
users_excel = APIRouter(prefix="/api/users/excel", tags=["Excel"], dependencies=[Depends(require_auth)])

ACTION_COLUMNS = [
    "SN", "Action", "Responsible", "Due", "Status", "Priority", "Section",
    "Plant", "Dept", "Machine", "Project", "Source", "Remarks", "Created",
]
USER_COLUMNS = [
    "Name", "Username", "Role", "Plant", "Dept", "Superior", "Phone",
    "Email", "Master", "Active", "Password",
]


def _sheet_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _hdr(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)


def _as_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    try:
        return _dt.date.fromisoformat(str(v).strip()[:10])
    except (ValueError, TypeError):
        return None


def _cell(v):
    if isinstance(v, _dt.datetime):
        return v.date()
    return v


# ── Actions export/import ──────────────────────────────────────────────────

@actions_excel.get("/export")
async def export_actions(
    plant_id: str = None, status: str = None, priority: str = None,
    db: AsyncSession = Depends(get_db), current_user: dict = Depends(get_current_user),
):
    from app.services.plant_scoping import scope_by_plant
    q = select(Action).where(Action.is_deleted.is_(False))
    if plant_id:
        q = q.where(Action.plant_id == plant_id)
    if status:
        q = q.where(Action.status == status)
    if priority:
        q = q.where(Action.priority == priority)
    q = scope_by_plant(q, current_user, Action.plant_id)
    q = q.order_by(Action.created.desc()).limit(5000)
    actions = (await db.execute(q)).scalars().all()

    plants = {p.id: p.name for p in (await db.execute(select(Plant))).scalars()}
    depts = {d.id: d.name for d in (await db.execute(select(Department))).scalars()}

    wb = Workbook()
    ws = wb.active
    ws.title = "Actions"
    _hdr(ws, ACTION_COLUMNS)
    for a in actions:
        ws.append([
            a.sn, a.text, a.responsible, _as_date(a.due), a.status, a.priority,
            a.section, plants.get(a.plant_id, ""), depts.get(a.dept_id, ""),
            a.machine_name, a.project_name, a.source, a.remarks, _as_date(a.created),
        ])
    return _sheet_response(wb, f"actions-{_dt.date.today().isoformat()}.xlsx")


async def _next_sn(db: AsyncSession) -> str:
    # mirror actions router SN format
    res = await db.execute(select(Action.sn).order_by(Action.created.desc()).limit(200))
    nums = []
    for (s,) in res:
        digits = "".join(ch for ch in (s or "") if ch.isdigit())
        if digits:
            nums.append(int(digits))
    return f"A{(max(nums) + 1) if nums else 1:04d}"


@actions_excel.post("/import")
async def import_actions(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db), current_user: dict = Depends(get_current_user),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Upload an .xlsx file")
    wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Empty sheet")
    headers = [str(h or "").strip().lower() for h in rows[0]]
    col = {h: i for i, h in enumerate(headers)}

    def get(r, *names):
        for n in names:
            if n.lower() in col:
                return _cell(r[col[n.lower()]])
        return None

    def name_row(r, *names):
        v = get(r, *names)
        return str(v).strip() if v not in (None, "") else None

    plants = {p.name: p.id for p in (await db.execute(select(Plant))).scalars() if p.name}
    depts = {d.name: d.id for d in (await db.execute(select(Department))).scalars() if d.name}

    created, errors = 0, []
    for idx, r in enumerate(rows[1:], start=2):
        if all(v in (None, "") for v in r):
            continue
        text = name_row(r, "action", "action point", "text", "task")
        if not text:
            errors.append(f"Row {idx}: missing 'Action' column value — skipped")
            continue
        payload = {
            "id": str(uuid.uuid4()),
            "text": text,
            "responsible": name_row(r, "responsible") or "",
            "due": _as_date(get(r, "due", "due date")),
            "status": (name_row(r, "status") or "NOT STARTED").upper(),
            "priority": (name_row(r, "priority") or "NORMAL").upper(),
            "section": name_row(r, "section") or "",
            "source": name_row(r, "source") or "Excel Import",
            "remarks": name_row(r, "remarks", "remark") or "",
            "src": "excel",
        }
        plant_name, dept_name = name_row(r, "plant"), name_row(r, "dept", "department")
        if plant_name:
            if plant_name in plants:
                payload["plant_id"] = plants[plant_name]
            else:
                errors.append(f"Row {idx}: unknown plant '{plant_name}' — plant left blank")
        if dept_name:
            if dept_name in depts:
                payload["dept_id"] = depts[dept_name]
            else:
                errors.append(f"Row {idx}: unknown dept '{dept_name}' — dept left blank")
        mach, proj = name_row(r, "machine", "machine name"), name_row(r, "project")
        if mach:
            payload["machine_name"] = mach
        if proj:
            payload["project_name"] = proj
        sn = name_row(r, "sn", "s.no", "sr no")
        if sn:
            dup = await db.execute(select(Action).where(Action.sn == sn))
            if dup.scalar_one_or_none():
                errors.append(f"Row {idx}: SN '{sn}' already exists — generated a new one")
            else:
                payload["sn"] = sn
        payload.setdefault("sn", await _next_sn(db))
        db.add(Action(**payload))
        created += 1
    await db.commit()
    return {"ok": True, "created": created, "errors": errors[:50]}


# ── Users export/import (admin only) ──────────────────────────────────────

def _require_admin(current_user: dict):
    if not current_user.get("is_admin"):
        raise HTTPException(403, "Admin access required")


@users_excel.get("/export")
async def export_users(db: AsyncSession = Depends(get_db), current_user: dict = Depends(get_current_user)):
    _require_admin(current_user)
    plants = {p.id: p.name for p in (await db.execute(select(Plant))).scalars()}
    depts = {d.id: d.name for d in (await db.execute(select(Department))).scalars()}
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    _hdr(ws, USER_COLUMNS)
    users = (await db.execute(select(User).where(User.is_deleted.is_(False)).order_by(User.name))).scalars().all()
    for u in users:
        ws.append([
            u.name, u.username, u.role, plants.get(u.plant_id, ""), depts.get(u.dept_id, ""),
            u.superior, u.phone, u.email,
            "YES" if u.master_access else "", "YES" if u.is_active else "", "",
        ])
    return _sheet_response(wb, f"users-{_dt.date.today().isoformat()}.xlsx")


@users_excel.post("/import")
async def import_users(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db), current_user: dict = Depends(get_current_user),
):
    _require_admin(current_user)
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Upload an .xlsx file")
    wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "Empty sheet")
    headers = [str(h or "").strip().lower() for h in rows[0]]
    col = {h: i for i, h in enumerate(headers)}

    def name_row(r, *names):
        for n in names:
            if n.lower() in col:
                v = _cell(r[col[n.lower()]])
                if v not in (None, ""):
                    return str(v).strip()
        return None

    plants = {p.name: p.id for p in (await db.execute(select(Plant))).scalars() if p.name}
    depts = {d.name: d.id for d in (await db.execute(select(Department))).scalars() if d.name}
    existing = {u.username for u in (await db.execute(select(User.username))).scalars()}

    created, updated, errors = 0, 0, []
    for idx, r in enumerate(rows[1:], start=2):
        if all(v in (None, "") for v in r):
            continue
        name = name_row(r, "name", "user", "full name")
        if not name:
            errors.append(f"Row {idx}: missing 'Name' — skipped")
            continue
        username = (name_row(r, "username") or name.lower().replace(" ", ".")).replace(" ", "")
        plant_name, dept_name = name_row(r, "plant"), name_row(r, "dept", "department")
        payload = {
            "id": str(uuid.uuid4()), "name": name, "username": username,
            "role": name_row(r, "role") or "User",
            "superior": name_row(r, "superior", "reports to"),
            "phone": name_row(r, "phone", "mobile"),
            "email": name_row(r, "email"),
            "master_access": str(name_row(r, "master", "master access") or "").upper() == "YES",
            "is_active": str(name_row(r, "active") or "YES").upper() != "NO",
        }
        if plant_name and plant_name in plants:
            payload["plant_id"] = plants[plant_name]
        elif plant_name:
            errors.append(f"Row {idx}: unknown plant '{plant_name}'")
        if dept_name and dept_name in depts:
            payload["dept_id"] = depts[dept_name]
        elif dept_name:
            errors.append(f"Row {idx}: unknown dept '{dept_name}'")
        password = name_row(r, "password")
        if password:
            payload["password"] = hash_password(password)
        if username in existing:
            errors.append(f"Row {idx}: username '{username}' already exists — skipped (use export+edit in Master Setup to update)")
            continue
        db.add(User(**payload))
        existing.add(username)
        created += 1
    await db.commit()
    return {"ok": True, "created": created, "updated": updated, "errors": errors[:50]}
